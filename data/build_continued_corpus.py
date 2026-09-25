"""Build the continued-pretraining corpus without reading data/raw.

The builder intentionally uses only:
- the already-processed original training corpus at data/processed/train.txt;
- data/MedQuad;
- data/Fine tuning 2.

It never scans data/ or data/raw recursively, so the original raw pretraining
sources cannot accidentally be re-ingested.
"""

import argparse
import csv
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from data.cleaner import clean_text
from data.dedup import is_near_duplicate, simhash
from data.file_hash import calculate_sha256
from data.filters import passes_basic_filters
from data.pdf_extractor import extract_pdf_text

DEFAULT_OLD_TRAIN = Path("data/processed/train.txt")
DEFAULT_SOURCES = (Path("data/MedQuad"), Path("data/Fine tuning 2"))
DEFAULT_OUTPUT = Path("data/processed/continued_pretraining")
SUPPORTED_TEXT = {".txt", ".md", ".markdown", ".text"}
SUPPORTED_STRUCTURED = {".json", ".jsonl", ".csv", ".tsv"}
SUPPORTED = SUPPORTED_TEXT | SUPPORTED_STRUCTURED | {".xml", ".pdf", ".xlsx"}

QUESTION_KEYS = {"question", "prompt", "instruction", "query", "user"}
ANSWER_KEYS = {"answer", "response", "output", "completion", "assistant"}
TEXT_KEYS = {"text", "content", "body"}
SKIP_KEYS = {"id", "source", "url", "metadata", "index", "uuid"}


def _content_hash(text: str) -> str:
    normalized = " ".join(text.split())
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _split_for_hash(content_hash: str) -> str:
    bucket = int(content_hash[:8], 16) % 100
    if bucket < 90:
        return "train"
    if bucket < 95:
        return "validation"
    return "test"


def _assert_allowed_path(path: Path) -> None:
    resolved = path.resolve()
    raw_dir = Path("data/raw").resolve()
    if resolved == raw_dir or raw_dir in resolved.parents:
        raise ValueError(f"data/raw is explicitly excluded: {path}")


def _normalise_record(record: object) -> str:
    if not isinstance(record, dict):
        return str(record) if record is not None else ""

    question_parts: list[str] = []
    answer_parts: list[str] = []
    text_parts: list[str] = []
    other_parts: list[str] = []

    for key, value in record.items():
        if value is None or key.lower() in SKIP_KEYS:
            continue
        if isinstance(value, (dict, list)):
            value_text = json.dumps(value, ensure_ascii=False)
        else:
            value_text = str(value).strip()
        if not value_text:
            continue
        key_lower = key.lower()
        if key_lower in QUESTION_KEYS:
            question_parts.append(value_text)
        elif key_lower in ANSWER_KEYS:
            answer_parts.append(value_text)
        elif key_lower in TEXT_KEYS:
            text_parts.append(value_text)
        else:
            other_parts.append(value_text)

    if question_parts or answer_parts:
        parts = []
        if question_parts:
            parts.append("Question: " + "\n".join(question_parts))
        if answer_parts:
            parts.append("Answer: " + "\n".join(answer_parts))
        return "\n".join(parts)
    if text_parts:
        return "\n".join(text_parts)
    return "\n".join(other_parts)


def _records_from_xml(path: Path) -> list[str]:
    root = ET.parse(path).getroot()
    records: list[str] = []
    pairs = list(root.iter("QAPair"))
    if pairs:
        for pair in pairs:
            question = (pair.findtext("Question") or "").strip()
            answer = (pair.findtext("Answer") or "").strip()
            if question and answer:
                records.append(f"Question: {question}\nAnswer: {answer}")
        return records
    text = " ".join(t.strip() for t in root.itertext() if t.strip())
    return [text] if text else []


def _records_from_json(path: Path) -> list[str]:
    if path.suffix.lower() == ".jsonl":
        records = []
        with path.open("r", encoding="utf-8", errors="replace") as handle:
            for line in handle:
                line = line.strip()
                if not line:
                    continue
                records.append(_normalise_record(json.loads(line)))
        return records
    payload = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    if isinstance(payload, list):
        return [_normalise_record(item) for item in payload]
    if isinstance(payload, dict):
        for key in ("data", "records", "examples", "items"):
            if isinstance(payload.get(key), list):
                return [_normalise_record(item) for item in payload[key]]
        return [_normalise_record(payload)]
    return [str(payload)]


def _records_from_delimited(path: Path) -> list[str]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    records: list[str] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=delimiter):
            records.append(_normalise_record(row))
    return records


def _records_from_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[str] = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        except StopIteration:
            continue
        for values in rows:
            record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
            records.append(_normalise_record(record))
    return records


def _extract_records(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".xml":
        return _records_from_xml(path)
    if suffix in {".json", ".jsonl"}:
        return _records_from_json(path)
    if suffix in {".csv", ".tsv"}:
        return _records_from_delimited(path)
    if suffix == ".xlsx":
        return _records_from_xlsx(path)
    if suffix == ".pdf":
        text = extract_pdf_text(path)
        return [text] if text.strip() else []
    if suffix in SUPPORTED_TEXT:
        text = path.read_text(encoding="utf-8", errors="replace")
        # Existing processed corpora and plain text often contain document
        # boundaries separated by blank lines. Preserve those as documents.
        return [part for part in re.split(r"\n\s*\n+", text) if part.strip()]
    return []


def _iter_source_files(root: Path):
    _assert_allowed_path(root)
    if not root.exists():
        raise FileNotFoundError(f"Source directory not found: {root}")
    if root.is_file():
        if root.suffix.lower() in SUPPORTED:
            yield root
        return
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            _assert_allowed_path(path)
            yield path


def _iter_clean_documents(sources: list[Path], old_train: Path):
    _assert_allowed_path(old_train)
    if not old_train.exists():
        raise FileNotFoundError(f"Original processed training corpus not found: {old_train}")

    # The old corpus is already cleaned and split. We use only train.txt here;
    # data/raw is never consulted.
    for index, text in enumerate(_extract_records(old_train)):
        yield f"old-train:{index}", clean_text(text)

    for source in sources:
        for path in _iter_source_files(source):
            for index, text in enumerate(_extract_records(path)):
                yield f"new:{path}:{index}", clean_text(text)


def build_corpus(sources: list[Path], old_train: Path, output_dir: Path) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_files = {
        split: (output_dir / f"{split}.txt").open("w", encoding="utf-8")
        for split in ("train", "validation", "test")
    }
    manifest_path = output_dir / "manifest.jsonl"
    stats = {
        "old_train_documents": 0,
        "source_files": 0,
        "accepted": 0,
        "rejected": 0,
        "duplicates": 0,
        "near_duplicates": 0,
        "extraction_failures": 0,
        "train_documents": 0,
        "validation_documents": 0,
        "test_documents": 0,
    }
    seen_hashes: set[str] = set()
    seen_simhashes: list[int] = []

    try:
        with manifest_path.open("w", encoding="utf-8") as manifest:
            for source_id, raw_text in _iter_clean_documents(sources, old_train):
                if source_id.startswith("old-train:"):
                    stats["old_train_documents"] += 1
                else:
                    if source_id.rsplit(":", 1)[-1] == "0":
                        stats["source_files"] += 1
                try:
                    cleaned = clean_text(raw_text)
                    if not passes_basic_filters(cleaned):
                        stats["rejected"] += 1
                        continue
                    content_hash = _content_hash(cleaned)
                    if content_hash in seen_hashes:
                        stats["duplicates"] += 1
                        continue
                    seen_hashes.add(content_hash)
                    fingerprint = simhash(cleaned)
                    if is_near_duplicate(fingerprint, seen_simhashes):
                        stats["near_duplicates"] += 1
                        continue
                    seen_simhashes.append(fingerprint)
                except Exception as exc:
                    stats["extraction_failures"] += 1
                    print(f"Failed to process {source_id}: {exc}")
                    continue

                # Keep the existing foundation corpus in training only. New
                # sources receive the deterministic 90/5/5 split.
                split = "train" if source_id.startswith("old-train:") else _split_for_hash(content_hash)
                output_files[split].write(cleaned + "\n\n")
                stats[f"{split}_documents"] += 1
                stats["accepted"] += 1
                manifest.write(
                    json.dumps(
                        {"source": source_id, "content_sha256": content_hash, "split": split},
                        ensure_ascii=False,
                    )
                    + "\n"
                )
    finally:
        for handle in output_files.values():
            handle.close()

    return stats


def main(args=None):
    parser = argparse.ArgumentParser(description="Build the continued-pretraining corpus.")
    parser.add_argument("--source", dest="sources", action="append", type=Path)
    parser.add_argument("--old-train", type=Path, default=DEFAULT_OLD_TRAIN)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parsed = parser.parse_args(args)
    sources = parsed.sources or list(DEFAULT_SOURCES)
    stats = build_corpus(sources, parsed.old_train, parsed.output_dir)
    print("\nContinued-pretraining corpus build complete")
    for key, value in stats.items():
        print(f"{key}: {value}")
    print(f"Output: {parsed.output_dir}")


if __name__ == "__main__":
    main()
