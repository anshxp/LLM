"""Ingest the complete local Fine tuning 2 corpus for second SFT.

Raw data stays local. The pipeline normalizes supervised records, audits malformed
and over-context examples, performs disk-backed global exact deduplication, and
writes deterministic JSONL shards. It never modifies the source corpus or trains.
"""
from __future__ import annotations
import argparse, csv, hashlib, json, re, sqlite3, zipfile
from collections import Counter
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree
from data.tokenizer import Tokenizer

SUPPORTED={".json",".jsonl",".csv",".tsv",".xml",".txt",".md",".parquet",".xlsx",".zip"}
QUESTION_FIELDS=("question","prompt","query","user","instruction")
ANSWER_FIELDS=("response","answer","output","completion","assistant")
ROLE_USER={"user","human","customer","question"}
ROLE_ASSISTANT={"assistant","bot","model","answer"}

def clean(value)->str:
    if value is None: return ""
    return re.sub(r"\s+"," ",str(value)).strip()

def source_name(path:str)->str: return path.replace("\\","/")

def make_record(instruction,input_text,response,category,source,source_id=""):
    instruction,input_text,response=map(clean,(instruction,input_text,response)); category=clean(category) or "fine_tuning_2"
    if not instruction or not response: return None
    record={"instruction":instruction,"input":input_text,"response":response,"category":category,"source":source_name(source),"source_id":clean(source_id)}
    record["id"]=hashlib.sha256(json.dumps(record,ensure_ascii=False,sort_keys=True).encode()).hexdigest(); return record

def first_value(obj,fields):
    if not isinstance(obj,dict): return ""
    lowered={str(k).lower():v for k,v in obj.items()}
    for field in fields:
        value=lowered.get(field.lower())
        if value is not None and clean(value): return value
    return ""

def conversation_record(obj,source,source_id):
    messages=obj if isinstance(obj,list) else obj.get("messages") if isinstance(obj,dict) else None
    if not isinstance(messages,list): return None
    user_parts=[]; assistant=""
    for message in messages:
        if not isinstance(message,dict): continue
        role=clean(message.get("role") or message.get("from") or message.get("speaker")).lower(); text=clean(message.get("content") or message.get("text") or message.get("value"))
        if not text: continue
        if role in ROLE_USER: user_parts.append(text)
        elif role in ROLE_ASSISTANT and user_parts: assistant=text
    if not user_parts or not assistant: return None
    return make_record("Answer the user's request.","\n\n".join(user_parts),assistant,"conversation",source,source_id)

def records_from_object(obj,source,source_id="")->Iterable[dict]:
    if isinstance(obj,dict):
        conv=conversation_record(obj,source,source_id)
        if conv: yield conv; return
        instruction=first_value(obj,("instruction",)); input_text=first_value(obj,("input","context")); response=first_value(obj,ANSWER_FIELDS)
        if instruction and response:
            record=make_record(instruction,input_text,response,first_value(obj,("category","type")),source,source_id)
            if record: yield record
            return
        question=first_value(obj,QUESTION_FIELDS); answer=first_value(obj,ANSWER_FIELDS)
        if question and answer:
            record=make_record("Answer the question accurately.",question,answer,"question_answer",source,source_id)
            if record: yield record
            return
        for key,value in obj.items(): yield from records_from_object(value,source,f"{source_id}.{key}" if source_id else str(key))
    elif isinstance(obj,list):
        for index,item in enumerate(obj): yield from records_from_object(item,source,f"{source_id}[{index}]")

def parse_json(text,source): yield from records_from_object(json.loads(text),source)
def parse_jsonl(text,source):
    for number,line in enumerate(text.splitlines(),1):
        if line.strip(): yield from records_from_object(json.loads(line),source,str(number))
def parse_delimited(text,source,delimiter):
    for number,row in enumerate(csv.DictReader(text.splitlines(),delimiter=delimiter),2): yield from records_from_object(dict(row),source,str(number))
def xml_text(element): return clean(" ".join(element.itertext())) if element is not None else ""
def parse_xml(text,source):
    root=ElementTree.fromstring(text)
    for node in root.iter():
        children={child.tag.split("}")[-1].lower():xml_text(child) for child in list(node)}; question=first_value(children,QUESTION_FIELDS); answer=first_value(children,ANSWER_FIELDS)
        if question and answer: yield make_record("Answer the question accurately.",question,answer,"question_answer",source,node.tag)
def parse_text(text,source):
    lines=[clean(x) for x in text.splitlines() if clean(x)]
    for i,line in enumerate(lines[:-1]):
        if line.endswith("?") and len(lines[i+1])>=20: yield make_record("Answer the question accurately.",line,lines[i+1],"question_answer",source,str(i+1))

def parse_parquet(path,source):
    import pyarrow.parquet as pq
    parquet=pq.ParquetFile(path)
    for batch in parquet.iter_batches(batch_size=10000):
        for index,row in enumerate(batch.to_pylist()): yield from records_from_object(row,source,str(index))

def parse_xlsx(path,source):
    from openpyxl import load_workbook
    workbook=load_workbook(path,read_only=True,data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows=sheet.iter_rows(values_only=True)
            try: headers=[clean(x).lower() for x in next(rows)]
            except StopIteration: continue
            for index,values in enumerate(rows,2):
                row={headers[i]:values[i] for i in range(min(len(headers),len(values))) if headers[i]}; yield from records_from_object(row,source,f"{sheet.title}:{index}")
    finally: workbook.close()

def iter_archive(path:Path):
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith("/"): continue
            suffix=Path(name).suffix.lower()
            if suffix not in SUPPORTED-{".zip",".parquet",".xlsx"}: continue
            yield source_name(f"{path}::{name}"),suffix,archive.read(name).decode("utf-8",errors="ignore")

def parse_source(suffix,text,source):
    if suffix==".json": yield from parse_json(text,source)
    elif suffix==".jsonl": yield from parse_jsonl(text,source)
    elif suffix in {".csv",".tsv"}: yield from parse_delimited(text,source,"\t" if suffix==".tsv" else ",")
    elif suffix==".xml": yield from parse_xml(text,source)
    elif suffix in {".txt",".md"}: yield from parse_text(text,source)

def iter_files(root:Path):
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED: yield path

def split_name(record_id:str)->str:
    value=int(record_id[:8],16)%100; return "train" if value<90 else "validation" if value<95 else "test"

def token_stats(record,tokenizer,context_length):
    prompt=f"### Instruction:\n{record['instruction']}\n\n### Input:\n{record['input']}\n\n### Response:\n"; prompt_ids=tokenizer.encode(prompt,add_bos=True); response_ids=tokenizer.encode(record["response"],add_eos=True); ids=prompt_ids+response_ids; unk_id=tokenizer.token_to_id["<unk>"]
    return {"prompt_tokens":len(prompt_ids),"response_tokens":len(response_ids),"total_tokens":len(ids),"unk_tokens":sum(x==unk_id for x in ids),"fits_context":len(ids)<=context_length+1}

def build(source_root:Path,output_dir:Path,tokenizer_path:Path,context_length:int,shard_size:int=50000):
    tokenizer=Tokenizer.from_file(tokenizer_path); output_dir.mkdir(parents=True,exist_ok=True)
    for old in output_dir.glob("*.jsonl"): old.unlink()
    for old in (output_dir/"manifest.json",output_dir/"dedup.sqlite3"): old.unlink(missing_ok=True)
    seen=sqlite3.connect(output_dir/"dedup.sqlite3"); seen.execute("PRAGMA journal_mode=WAL"); seen.execute("CREATE TABLE seen (key BLOB PRIMARY KEY)"); seen.commit()
    counts=Counter(); sources=Counter(); rejects=Counter(); lengths=Counter(); total_tokens=unk_tokens=0; handles={}; shard_counts=Counter(); rejection_handle=(output_dir/"rejections.jsonl").open("w",encoding="utf-8")
    def handle_for(split):
        index=shard_counts[split]//shard_size; key=(split,index)
        if key not in handles: handles[key]=(output_dir/f"{split}-{index:05d}.jsonl").open("w",encoding="utf-8")
        return handles[key]
    try:
        for path in iter_files(source_root):
            relative=source_name(str(path.relative_to(source_root)))
            try:
                suffix=path.suffix.lower()
                if suffix in {".parquet",".xlsx"}: sources_iter=[(relative,suffix,path)]
                elif suffix==".zip": sources_iter=list(iter_archive(path))
                else: sources_iter=[(relative,suffix,path.read_text(encoding="utf-8",errors="ignore"))]
                for source,suffix,payload in sources_iter:
                    counts["files_or_members"]+=1
                    try:
                        if suffix==".parquet": records=parse_parquet(payload,source)
                        elif suffix==".xlsx": records=parse_xlsx(payload,source)
                        else: records=parse_source(suffix,payload,source)
                        for record in records:
                            if record is None: rejects["empty_or_invalid"]+=1; continue
                            key_text="\n".join((record["instruction"].lower(),record["input"].lower(),record["response"].lower())); key_hash=hashlib.sha256(key_text.encode()).digest()
                            inserted=seen.execute("INSERT OR IGNORE INTO seen(key) VALUES (?)",(key_hash,)).rowcount
                            if not inserted: rejects["duplicate"]+=1; continue
                            stats=token_stats(record,tokenizer,context_length); total_tokens+=stats["total_tokens"]; unk_tokens+=stats["unk_tokens"]; lengths["within_context" if stats["fits_context"] else "over_context"]+=1
                            if not stats["fits_context"]:
                                rejects["over_context"]+=1; rejection_handle.write(json.dumps({"reason":"over_context","record":record,"stats":stats},ensure_ascii=False)+"\n"); continue
                            split=split_name(record["id"]); record.update(stats); handle_for(split).write(json.dumps(record,ensure_ascii=False)+"\n"); shard_counts[split]+=1; sources[source]+=1
                        seen.commit()
                    except Exception as exc:
                        rejects[f"parse_error:{type(exc).__name__}"]+=1; rejection_handle.write(json.dumps({"reason":f"parse_error:{type(exc).__name__}","source":source,"error":str(exc)},ensure_ascii=False)+"\n")
            except Exception as exc:
                rejects[f"file_error:{type(exc).__name__}"]+=1; rejection_handle.write(json.dumps({"reason":f"file_error:{type(exc).__name__}","source":relative,"error":str(exc)},ensure_ascii=False)+"\n")
    finally:
        rejection_handle.close(); seen.commit(); seen.close()
        for handle in handles.values(): handle.close()
    accepted=sum(shard_counts.values()); manifest={"source_root":str(source_root),"tokenizer":str(tokenizer_path),"context_length":context_length,"files_or_archive_members_seen":counts["files_or_members"],"unique_records_seen":sum(shard_counts.values())+rejects["over_context"],"accepted_records":accepted,"train_records":shard_counts["train"],"validation_records":shard_counts["validation"],"test_records":shard_counts["test"],"rejections":dict(rejects),"length_distribution":dict(lengths),"average_tokens":total_tokens/max(1,sum(lengths.values())),"unk_token_rate":unk_tokens/max(1,total_tokens),"source_counts":dict(sources),"split_policy":"stable hash of record id: 90% train, 5% validation, 5% test","deduplication":"disk-backed SQLite SHA-256 key index"}
    (output_dir/"manifest.json").write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding="utf-8"); return manifest

def main():
    parser=argparse.ArgumentParser(description="Prepare Fine tuning 2 for second SFT."); parser.add_argument("--source-root",type=Path,default=Path("data/Fine tuning 2")); parser.add_argument("--output-dir",type=Path,default=Path("data/instruction_v2")); parser.add_argument("--tokenizer",type=Path,default=Path("data/processed/tokenizer.json")); parser.add_argument("--context-length",type=int,default=256); parser.add_argument("--shard-size",type=int,default=50000); args=parser.parse_args()
    if not args.source_root.is_dir(): raise FileNotFoundError(f"Fine tuning 2 directory not found: {args.source_root}")
    if args.context_length<2 or args.shard_size<=0: raise ValueError("context-length must be >= 2 and shard-size must be positive")
    print(json.dumps(build(args.source_root,args.output_dir,args.tokenizer,args.context_length,args.shard_size),indent=2,ensure_ascii=False))
if __name__=="__main__": main()
